"""Spreadsheet parsing and M4-compatible multi-source acquisition."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
from urllib.parse import unquote

from pydantic import BaseModel, Field

import award_audit.agent.toolkit.web as web
from award_audit.agent.toolkit.safety import inspect_evidence_file

MAX_EXCEL_ROWS_PER_SHEET = 100_000
SheetRows = tuple[str, list[list[str]], bool]

_ROLE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "team": ("团队名称", "队伍名称", "参赛队伍名称", "参赛队伍"),
    "organization": ("参赛单位", "组织单位", "获奖单位"),
    "instructor_or_person": ("指导教师姓名", "指导导师姓名", "获奖人姓名", "姓名"),
    "work_or_project": ("项目名称", "作品名称", "成果名称", "专利名称"),
}
_CATEGORY_HEADER_ALIASES = ("组别", "赛道", "赛道类别", "类别", "项目类别", "奖项类别")
_LEVEL_HEADER_ALIASES = ("奖项等级", "获奖等级", "等级", "奖项")


class SpreadsheetRowLimitError(RuntimeError):
    """Raised when a submitted workbook cannot be read completely."""


def _assemble_sheets(sheets: list[SheetRows], *, row_limit: int) -> dict[str, object]:
    truncated_sheets = [title for title, _rows, truncated in sheets if truncated]
    metadata: dict[str, object] = {
        "truncated": bool(truncated_sheets),
        "truncated_sheets": truncated_sheets,
        "row_limit": row_limit,
        "sheet_row_counts": {title: len(rows) for title, rows, _truncated in sheets},
        "sheet_grids": [
            {
                "sheet": title,
                "n_rows": len(rows),
                "rows": rows,
                "truncated": truncated,
            }
            for title, rows, truncated in sheets
        ],
    }
    if not sheets:
        return {"sheet": "", "n_rows": 0, "rows": [], **metadata}
    if len(sheets) == 1:
        title, rows, _truncated = sheets[0]
        return {"sheet": title, "n_rows": len(rows), "rows": rows, **metadata}
    merged: list[list[str]] = []
    for title, rows, _truncated in sheets:
        merged.append([f"【等级：{title}】"])
        merged.extend(rows)
    return {
        "sheet": " / ".join(title for title, _rows, _truncated in sheets),
        "sheets": [title for title, _rows, _truncated in sheets],
        "n_rows": len(merged),
        "rows": merged,
        **metadata,
    }


def _read_xlsx_sheets(path: Path, max_rows: int) -> list[SheetRows]:
    import openpyxl

    sheets: list[SheetRows] = []
    with path.open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows: list[list[str]] = []
                truncated = False
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if value is None else str(value).strip() for value in row]
                    if any(cells):
                        if len(rows) >= max_rows:
                            truncated = True
                            break
                        rows.append(cells)
                if len(rows) >= 2 or truncated:
                    sheets.append((sheet.title, rows, truncated))
        finally:
            workbook.close()
    return sheets


def _xls_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xls_sheets(path: Path, max_rows: int) -> list[SheetRows]:
    try:
        import xlrd  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RuntimeError("读取老 .xls 需要 xlrd：pip install xlrd") from exc
    book = xlrd.open_workbook(str(path))
    sheets: list[SheetRows] = []
    for sheet in book.sheets():
        rows: list[list[str]] = []
        truncated = False
        for row_index in range(sheet.nrows):
            cells = [_xls_cell(sheet.cell_value(row_index, column_index))
                     for column_index in range(sheet.ncols)]
            if any(cells):
                if len(rows) >= max_rows:
                    truncated = True
                    break
                rows.append(cells)
        if len(rows) >= 2 or truncated:
            sheets.append((sheet.name, rows, truncated))
    return sheets


def parse_award_excel(
    path: Path,
    max_rows: int = MAX_EXCEL_ROWS_PER_SHEET,
) -> dict[str, object]:
    """Parse a validated XLS/XLSX by magic, preserving M4's grid shape."""

    inspection = inspect_evidence_file(path, max_bytes=20 * 1024 * 1024,
                                       allowed_kinds={"xlsx", "xls"})
    if inspection.kind == "xls":
        sheets = _read_xls_sheets(path, max_rows)
    else:
        sheets = _read_xlsx_sheets(path, max_rows)
    return _assemble_sheets(sheets, row_limit=max_rows)


def extract_semantic_roster_records(grid: dict[str, object]) -> list[dict[str, object]]:
    """Extract table-role identities without treating auxiliary columns as awards."""

    raw_sheets = grid.get("sheet_grids", [])
    sheets = raw_sheets if isinstance(raw_sheets, list) else []
    records: list[dict[str, object]] = []
    for raw_sheet in sheets:
        if not isinstance(raw_sheet, dict):
            continue
        sheet_name = str(raw_sheet.get("sheet", ""))
        raw_rows = raw_sheet.get("rows", [])
        rows = raw_rows if isinstance(raw_rows, list) else []
        header_index = -1
        role_type = ""
        identity_column = -1
        headers: list[str] = []
        for index, raw_row in enumerate(rows[:30]):
            if not isinstance(raw_row, list):
                continue
            candidate = [str(cell or "").strip() for cell in raw_row]
            candidate_roles = [
                (role, column_index)
                for role, aliases in _ROLE_HEADER_ALIASES.items()
                for column_index, value in enumerate(candidate)
                if value in aliases
            ]
            if not candidate_roles:
                continue
            title_text = " ".join(
                str(cell or "")
                for previous in rows[:index]
                if isinstance(previous, list)
                for cell in previous
            )
            if "组织" in title_text and any(role == "organization" for role, _ in candidate_roles):
                selected_role = "organization"
            elif any(role == "team" for role, _ in candidate_roles):
                selected_role = "team"
            elif any(role == "work_or_project" for role, _ in candidate_roles):
                selected_role = "work_or_project"
            elif (
                any(term in title_text for term in ("教师", "导师", "个人"))
                and any(role == "instructor_or_person" for role, _ in candidate_roles)
            ):
                selected_role = "instructor_or_person"
            elif len({role for role, _ in candidate_roles}) == 1:
                selected_role = candidate_roles[0][0]
            else:
                continue
            header_index = index
            role_type = selected_role
            identity_column = next(
                column for role, column in candidate_roles if role == selected_role
            )
            headers = candidate
            break
        if header_index < 0:
            continue
        title = " ".join(
            str(cell or "").strip()
            for row in rows[:header_index]
            if isinstance(row, list)
            for cell in row
            if str(cell or "").strip()
        )
        category_columns = [
            index for index, value in enumerate(headers)
            if value in _CATEGORY_HEADER_ALIASES
        ]
        level_columns = [
            index for index, value in enumerate(headers)
            if value in _LEVEL_HEADER_ALIASES
        ]
        for row_index, raw_row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if not isinstance(raw_row, list) or identity_column >= len(raw_row):
                continue
            identity = str(raw_row[identity_column] or "").strip()
            if not identity or identity in _ROLE_HEADER_ALIASES[role_type]:
                continue
            category_values = [
                str(raw_row[index] or "").strip()
                for index in category_columns
                if index < len(raw_row) and str(raw_row[index] or "").strip()
            ]
            level_values = [
                str(raw_row[index] or "").strip()
                for index in level_columns
                if index < len(raw_row) and str(raw_row[index] or "").strip()
            ]
            records.append({
                "sheet": sheet_name,
                "row_number": row_index,
                "role_type": role_type,
                "identity": identity,
                "identity_field": headers[identity_column],
                "row_values": {
                    header: str(raw_row[column] or "").strip()
                    for column, header in enumerate(headers)
                    if header and column < len(raw_row)
                },
                "title": title,
                "category_values": category_values,
                "level_values": level_values,
                "document_complete": not bool(raw_sheet.get("truncated", False)),
            })
    return records


class AcquiredDocument(BaseModel):
    label: str
    source_url: str
    page_url: str
    raw_path: Path
    grid: dict[str, object]
    page_year: str = ""


class AcquiredGrid(BaseModel):
    source_url: str
    source_urls: list[str] = Field(default_factory=list)
    page_url: str
    raw_path: Path
    raw_paths: list[Path] = Field(default_factory=list)
    grid: dict[str, object]
    found_assets: list[str] = Field(default_factory=list)
    documents: list[AcquiredDocument] = Field(default_factory=list)
    discovered_attachment_urls: list[str] = Field(default_factory=list)
    attempted_attachment_urls: list[str] = Field(default_factory=list)
    failed_attachment_urls: list[str] = Field(default_factory=list)
    unprocessed_attachment_urls: list[str] = Field(default_factory=list)
    attachment_parent_urls: dict[str, str] = Field(default_factory=dict)
    attachment_errors: dict[str, str] = Field(default_factory=dict)
    all_attachments_processed: bool = True


def _source_label(attachment_text: str, page_year: str) -> str:
    if page_year and not web.extract_years(attachment_text):
        return f"{attachment_text}（{page_year}）"
    return attachment_text


FetchPage = Callable[[str, float], web.PageContent]
DownloadFile = Callable[..., Path]
ParseExcel = Callable[[Path, int], dict[str, object]]


def acquire_excel_grid(
    urls: list[str],
    workdir: Path,
    max_total: int = 100,
    *,
    direct_attachment_urls: list[str] | None = None,
    direct_referer: str = "",
    direct_attachment_parent_urls: dict[str, str] | None = None,
    fetch_page_fn: Callable[..., web.PageContent] | None = None,
    download_file_fn: DownloadFile | None = None,
    parse_excel_fn: Callable[..., dict[str, object]] | None = None,
    attachment_filter_fn: Callable[[web.Attachment], bool] | None = None,
    attachment_error_fn: Callable[[web.Attachment, Exception], None] | None = None,
) -> AcquiredGrid | None:
    """Acquire and combine all parseable Excel attachments from known pages."""

    fetcher = fetch_page_fn or web.fetch_page
    downloader = download_file_fn or web.download_file
    parser = parse_excel_fn or parse_award_excel
    assets: list[str] = []
    parts: list[tuple[str, str, str, Path, dict[str, object], str]] = []
    attempted_urls: set[str] = set()
    attempted_order: list[str] = []
    discovered_attachment_urls: list[str] = []
    attachment_parent_urls: dict[str, str] = {}
    attachment_errors: dict[str, str] = {}
    page_hit = ""

    def acquire_attachment(
        attachment: web.Attachment,
        *,
        referer: str,
        page_year: str = "",
    ) -> None:
        nonlocal page_hit
        if attachment.url in attempted_urls:
            return
        if attachment_filter_fn is not None and not attachment_filter_fn(attachment):
            return
        if len(attempted_urls) >= max_total:
            return
        attempted_urls.add(attachment.url)
        attempted_order.append(attachment.url)
        last_error: Exception | None = None
        for _attempt in range(2):
            try:
                local = downloader(
                    attachment.url,
                    workdir,
                    excel_only=True,
                    referer=referer,
                )
                grid = parser(local)
                break
            except Exception as exc:  # noqa: BLE001 - retain terminal failure for M5.
                last_error = exc
        else:
            assert last_error is not None
            attachment_errors[attachment.url] = (
                f"{type(last_error).__name__}: {str(last_error)[:500]}"
            )
            if attachment_error_fn is not None:
                attachment_error_fn(attachment, last_error)
            return
        rows = grid.get("rows")
        if isinstance(rows, list) and rows:
            if not page_hit:
                page_hit = referer
            parts.append(
                (
                    attachment.text,
                    attachment.url,
                    referer,
                    local,
                    grid,
                    page_year,
                )
            )

    for url in urls:
        try:
            page = fetcher(url)
        except Exception:  # noqa: BLE001 - one source must not kill the batch
            continue
        if page.status != 200:
            continue
        for attachment in page.attachments:
            if attachment.url not in assets:
                assets.append(attachment.url)
            if attachment.url not in discovered_attachment_urls:
                discovered_attachment_urls.append(attachment.url)
            attachment_parent_urls.setdefault(attachment.url, page.url or url)
        for image in page.images:
            if image not in assets:
                assets.append(image)
        title_years = web.extract_years(page.title)
        page_year = next(iter(title_years)) if len(title_years) == 1 else ""
        for attachment in page.attachments:
            acquire_attachment(attachment, referer=page.url or url, page_year=page_year)
    referer = direct_referer or (urls[0] if urls else "")
    for attachment_url in dict.fromkeys(direct_attachment_urls or []):
        if attachment_url not in assets:
            assets.append(attachment_url)
        if attachment_url not in discovered_attachment_urls:
            discovered_attachment_urls.append(attachment_url)
        label = unquote(attachment_url.split("?", 1)[0]).rsplit("/", 1)[-1]
        attachment_parent_urls.setdefault(
            attachment_url,
            (direct_attachment_parent_urls or {}).get(attachment_url, referer),
        )
        acquire_attachment(
            web.Attachment(
                text=label or "bound M4 attachment",
                url=attachment_url,
                is_excel=Path(label).suffix.casefold() in web.EXCEL_EXTS,
            ),
            referer=(direct_attachment_parent_urls or {}).get(attachment_url, referer),
        )
    if not parts:
        return None
    failed_attachment_urls = [
        url for url in discovered_attachment_urls if url in attachment_errors
    ]
    unprocessed_attachment_urls = [
        url for url in discovered_attachment_urls if url not in attempted_urls
    ]
    all_attachments_processed = not (
        failed_attachment_urls or unprocessed_attachment_urls
    )
    documents = [
        AcquiredDocument(
            label=text,
            source_url=attachment_url,
            page_url=page_url,
            raw_path=local,
            grid=grid,
            page_year=page_year,
        )
        for text, attachment_url, page_url, local, grid, page_year in parts
    ]
    if len(parts) == 1:
        _text, attachment_url, _page_url, local, grid, _page_year = parts[0]
        return AcquiredGrid(source_url=attachment_url, source_urls=[attachment_url],
                            page_url=page_hit, raw_path=local, raw_paths=[local], grid=grid,
                            found_assets=assets, documents=documents,
                            discovered_attachment_urls=discovered_attachment_urls,
                            attempted_attachment_urls=attempted_order,
                            failed_attachment_urls=failed_attachment_urls,
                            unprocessed_attachment_urls=unprocessed_attachment_urls,
                            attachment_parent_urls=attachment_parent_urls,
                            attachment_errors=attachment_errors,
                            all_attachments_processed=all_attachments_processed)
    merged: list[list[str]] = []
    for text, _url, _page_url, _local, grid, page_year in parts:
        merged.append([f"【名单：{_source_label(text, page_year)}】"])
        rows = grid.get("rows")
        if isinstance(rows, list):
            merged.extend(rows)
    combined: dict[str, object] = {
        "sheet": " / ".join(text for text, _, _, _, _, _ in parts),
        "sheets": [text for text, _, _, _, _, _ in parts],
        "n_rows": len(merged),
        "rows": merged,
    }
    return AcquiredGrid(source_url=parts[0][1], source_urls=[part[1] for part in parts],
                        page_url=page_hit, raw_path=parts[0][3],
                        raw_paths=[part[3] for part in parts], grid=combined,
                        found_assets=assets, documents=documents,
                        discovered_attachment_urls=discovered_attachment_urls,
                        attempted_attachment_urls=attempted_order,
                        failed_attachment_urls=failed_attachment_urls,
                        unprocessed_attachment_urls=unprocessed_attachment_urls,
                        attachment_parent_urls=attachment_parent_urls,
                        attachment_errors=attachment_errors,
                        all_attachments_processed=all_attachments_processed)
